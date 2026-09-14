from __future__ import annotations

import argparse
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path
import sys

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt

PACKAGE_ROOT = Path(__file__).resolve().parent
sys.path.insert(0, str(PACKAGE_ROOT))

from activity_diagram_to_usecase.styles import (  # noqa: E402
    HEADER_FILL,
    SECTION_FILL,
    configure_document,
    configure_table,
    set_cell_margins,
    set_column_widths,
    set_repeat_table_header,
    shade_cell,
    style_paragraph,
    style_run,
)


@dataclass
class TextRun:
    text: str
    bold: bool = False
    italic: bool = False


@dataclass
class CellBlock:
    runs: list[TextRun] = field(default_factory=list)
    bullet: bool = False


@dataclass
class HtmlCell:
    blocks: list[CellBlock] = field(default_factory=list)
    colspan: int = 1
    header: bool = False


class TableParser(HTMLParser):
    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[HtmlCell]] = []
        self.row: list[HtmlCell] | None = None
        self.cell: HtmlCell | None = None
        self.block: CellBlock | None = None
        self.styles: list[str] = []

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        tag = tag.lower()
        attributes = dict(attrs)
        if tag == "tr":
            self.row = []
        elif tag in {"td", "th"} and self.row is not None:
            colspan = int(attributes.get("colspan") or "1")
            self.cell = HtmlCell(colspan=max(colspan, 1), header=tag == "th")
            self.block = CellBlock()
        elif tag == "li" and self.cell is not None:
            self._finish_block()
            self.block = CellBlock(bullet=True)
            self.block.runs.append(TextRun("• "))
        elif tag == "br" and self.cell is not None:
            self._append_text("\n")
        elif tag in {"b", "strong", "i", "em"}:
            self.styles.append(tag)

    def handle_endtag(self, tag: str) -> None:
        tag = tag.lower()
        if tag in {"b", "strong", "i", "em"}:
            for index in range(len(self.styles) - 1, -1, -1):
                if self.styles[index] == tag:
                    self.styles.pop(index)
                    break
        elif tag in {"td", "th"} and self.cell is not None and self.row is not None:
            self._finish_block()
            self.row.append(self.cell)
            self.cell = None
            self.block = None
        elif tag == "tr" and self.row is not None:
            if self.row:
                self.rows.append(self.row)
            self.row = None

    def handle_data(self, data: str) -> None:
        if self.cell is not None:
            self._append_text(data)

    def _append_text(self, data: str) -> None:
        data = data.replace("\xa0", " ")
        data = re.sub(r"[ \t\r\n]+", " ", data)
        if not data:
            return
        if self.block is None:
            self.block = CellBlock()
        bold = any(tag in {"b", "strong"} for tag in self.styles)
        italic = any(tag in {"i", "em"} for tag in self.styles)
        self.block.runs.append(TextRun(data, bold=bold, italic=italic))

    def _finish_block(self) -> None:
        if self.block is not None and any(run.text for run in self.block.runs):
            self.cell.blocks.append(self.block)
        self.block = None


def parse_markdown(path: Path) -> tuple[str, list[list[HtmlCell]]]:
    text = path.read_text(encoding="utf-8")
    heading = re.search(r"^#\s+(.+?)\s*$", text, flags=re.MULTILINE)
    title = heading.group(1).strip() if heading else path.stem.replace("-", " ").title()
    parser = TableParser()
    table_match = re.search(r"<table\b[^>]*>(.*?)</table>", text, flags=re.IGNORECASE | re.DOTALL)
    if table_match is None:
        raise ValueError(f"No HTML table found in {path}")
    parser.feed(table_match.group(0))
    if not parser.rows:
        raise ValueError(f"HTML table is empty in {path}")
    return title, parser.rows


def build_document(title: str, rows: list[list[HtmlCell]]) -> Document:
    document = Document()
    configure_document(document)
    title_paragraph = document.add_paragraph(style="Title")
    title_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_paragraph.paragraph_format.space_after = Pt(10)
    title_run = title_paragraph.add_run(title)
    style_run(title_run, bold=True, color="000000")

    table = document.add_table(rows=0, cols=2)
    configure_table(table)
    for source_row in rows:
        row = table.add_row()
        cells = _normalise_cells(source_row)
        if len(cells) == 1:
            target = row.cells[0].merge(row.cells[1])
            _render_cell(target, cells[0])
            if _is_section(cells[0]):
                shade_cell(target, SECTION_FILL)
        else:
            _render_cell(row.cells[0], cells[0])
            _render_cell(row.cells[1], cells[1])
            if cells[0].header and cells[1].header:
                for cell in row.cells:
                    shade_cell(cell, HEADER_FILL)
                    for paragraph in cell.paragraphs:
                        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        for run in paragraph.runs:
                            style_run(run, bold=True, color="FFFFFF")
                set_repeat_table_header(row)

    set_column_widths(table)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)
    return document


def _normalise_cells(source_row: list[HtmlCell]) -> list[HtmlCell]:
    if len(source_row) == 1 or any(cell.colspan >= 2 for cell in source_row):
        return [source_row[0]]
    if len(source_row) >= 2:
        return source_row[:2]
    return [HtmlCell()]


def _render_cell(cell, source: HtmlCell) -> None:
    cell.text = ""
    for index, block in enumerate(source.blocks):
        paragraph = cell.paragraphs[0] if index == 0 else cell.add_paragraph()
        style_paragraph(paragraph)
        for run_data in block.runs:
            run = paragraph.add_run(run_data.text)
            style_run(run, bold=run_data.bold, italic=run_data.italic)


def _is_section(cell: HtmlCell) -> bool:
    text = "".join(run.text for block in cell.blocks for run in block.runs).strip()
    return bool(text.endswith(":")) and all(
        run.bold for block in cell.blocks for run in block.runs if run.text.strip()
    )


def convert_file(input_path: Path, output_path: Path) -> Path:
    title, rows = parse_markdown(input_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    build_document(title, rows).save(output_path)
    return output_path


def markdown_files(input_path: Path) -> list[Path]:
    if input_path.is_file():
        return [input_path]
    return sorted(
        path for path in input_path.rglob("*.md") if path.name.lower() != "readme.md"
    )


def _output_path(input_path: Path, output: Path | None) -> Path:
    output = output or input_path.with_suffix(".docx")
    if output.suffix.lower() != ".docx":
        output = output / input_path.with_suffix(".docx").name
    return output


def _browse_folder(title: str) -> Path | None:
    try:
        import tkinter as tk
        from tkinter import filedialog

        root = tk.Tk()
        root.withdraw()
        root.attributes("-topmost", True)
        selected = filedialog.askdirectory(title=title, initialdir=str(Path.cwd()))
        root.destroy()
        return Path(selected) if selected else None
    except Exception as error:
        print(f"Folder picker unavailable: {error}")
        return None


def run_console() -> int:
    print("=== Use-case Markdown to DOCX ===")
    print("1. Convert one Markdown file")
    print("2. Convert all Markdown files in a folder")
    choice = input("Choose an option [2]: ").strip() or "2"
    if choice not in {"1", "2"}:
        print("Invalid option.")
        return 1

    if choice == "2":
        print("Select the Markdown folder in the folder picker.")
        source = _browse_folder("Select Markdown folder")
        if source is None:
            manual_source = input("Enter the folder path manually, or press Enter to cancel: ").strip()
            if not manual_source:
                print("No folder selected.")
                return 1
            source = Path(manual_source.strip('"')).expanduser()
    else:
        source = Path(input("Enter the Markdown file path: ").strip().strip('"')).expanduser()
    if not source.exists():
        print(f"Input does not exist: {source}")
        return 1

    if choice == "1":
        if not source.is_file() or source.suffix.lower() != ".md":
            print("Please select a Markdown file.")
            return 1
        output = input("Output folder or DOCX path [same folder]: ").strip().strip('"')
        target = _output_path(source, Path(output) if output else None)
        print(f"Created: {convert_file(source, target)}")
        return 0

    if not source.is_dir():
        print("Please select a folder.")
        return 1
    output = input("Output folder [use-case-docx]: ").strip().strip('"')
    output_dir = Path(output).expanduser() if output else Path("use-case-docx")
    files = markdown_files(source)
    if not files:
        print(f"No Markdown files found in: {source}")
        return 1
    for path in files:
        relative = path.relative_to(source).with_suffix(".docx")
        print(f"Created: {convert_file(path, output_dir / relative)}")
    print(f"Completed: {len(files)} file(s)")
    return 0


def main() -> int:
    parser = argparse.ArgumentParser(description="Convert use-case Markdown tables to DOCX.")
    parser.add_argument(
        "input",
        nargs="?",
        type=Path,
        help="A single Markdown file",
    )
    parser.add_argument(
        "--folder",
        type=Path,
        help="Recursively convert every Markdown file in this folder",
    )
    parser.add_argument(
        "--output",
        type=Path,
        help="DOCX path for one file, or output directory for a directory input",
    )
    parser.add_argument(
        "--merged",
        action="store_true",
        help="Merge all Markdown files from a directory into one DOCX",
    )
    parser.add_argument(
        "--interactive",
        "--console",
        dest="interactive",
        action="store_true",
        help="Open the interactive console menu",
    )
    args = parser.parse_args()
    if args.interactive or (args.input is None and args.folder is None):
        return run_console()
    if args.input is not None and args.folder is not None:
        parser.error("Use either a Markdown file or --folder, not both")
    input_arg = args.folder or args.input
    if input_arg is None:
        parser.error("Provide a Markdown file or --folder")
    input_path = input_arg.resolve()
    if not input_path.exists():
        parser.error(f"Input does not exist: {input_path}")
    files = markdown_files(input_path)
    if not files:
        parser.error(f"No Markdown files found in: {input_path}")

    if args.merged:
        if input_path.is_file():
            parser.error("--merged requires a directory input")
        output = args.output or Path(f"{input_path.name}.docx")
        if output.suffix.lower() != ".docx":
            output = output / f"{input_path.name}.docx"
        first_title, first_rows = parse_markdown(files[0])
        document = build_document(first_title, first_rows)
        for path in files[1:]:
            document.add_page_break()
            title, rows = parse_markdown(path)
            source_document = build_document(title, rows)
            for element in source_document.element.body:
                if element.tag.endswith("sectPr"):
                    continue
                document.element.body.append(element)
        output.parent.mkdir(parents=True, exist_ok=True)
        document.save(output)
        print(output)
        return 0

    if input_path.is_file():
        print(convert_file(input_path, _output_path(input_path, args.output)))
        return 0

    output_dir = args.output or input_path.parent / f"{input_path.name}-docx"
    for path in files:
        relative = path.relative_to(input_path).with_suffix(".docx")
        print(convert_file(path, output_dir / relative))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
