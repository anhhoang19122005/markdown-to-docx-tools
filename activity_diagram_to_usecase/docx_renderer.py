from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches

from .models import FlowItem, FlowSection, UseCase
from .styles import (
    HEADER_FILL,
    SECTION_FILL,
    configure_document,
    configure_table,
    set_cell_margins,
    set_column_widths,
    set_repeat_table_header,
    shade_cell,
    style_paragraph,
    style_run,
)


def render_use_case(use_case: UseCase, output_path: str | Path) -> Path:
    document = Document()
    configure_document(document)
    _append_use_case_table(document, use_case)
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def render_merged(use_cases: list[UseCase], output_path: str | Path) -> Path:
    document = Document()
    configure_document(document)
    for index, use_case in enumerate(use_cases):
        if index:
            document.add_page_break()
        _append_use_case_table(document, use_case)
    path = Path(output_path)
    path.parent.mkdir(parents=True, exist_ok=True)
    document.save(path)
    return path


def _append_use_case_table(document, use_case: UseCase) -> None:
    table = document.add_table(rows=0, cols=2)
    table.style = "Table Grid"
    configure_table(table)

    _merged_text_row(table, "Use case name:", use_case.name)
    _merged_text_row(table, "Brief description:", use_case.description)
    _merged_text_row(table, "Primary actor:", use_case.primary_actor)
    _merged_text_row(
        table,
        "Secondary actors:",
        ", ".join(use_case.secondary_actors) if use_case.secondary_actors else "None",
    )
    _condition_row(table, "Precondition:", use_case.preconditions)
    _condition_row(table, "Postcondition:", use_case.postconditions)

    _merged_section_row(table, "Main flow:")
    header = table.add_row()
    _cell_text(
        header.cells[0],
        "Actor",
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        color="FFFFFF",
    )
    _cell_text(
        header.cells[1],
        "System",
        bold=True,
        alignment=WD_ALIGN_PARAGRAPH.CENTER,
        color="FFFFFF",
    )
    shade_cell(header.cells[0], HEADER_FILL)
    shade_cell(header.cells[1], HEADER_FILL)
    set_repeat_table_header(header)
    _flow_rows(table, use_case.main_flow)

    _flow_section(table, "Alternate flow:", use_case.alternate_flows)
    _flow_section(table, "Exception flow:", use_case.exception_flows)

    set_column_widths(table)
    for row in table.rows:
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_margins(cell)


def _merged_text_row(table, label: str, value: str) -> None:
    row = table.add_row()
    cell = row.cells[0].merge(row.cells[1])
    _cell_text(cell, "", clear=True)
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph)
    label_run = paragraph.add_run(label + " ")
    style_run(label_run, bold=True)
    if value:
        value_run = paragraph.add_run(value)
        style_run(value_run)


def _condition_row(table, label: str, values: list[str]) -> None:
    row = table.add_row()
    cell = row.cells[0].merge(row.cells[1])
    _cell_text(cell, "", clear=True)
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph)
    label_run = paragraph.add_run(label)
    style_run(label_run, bold=True)
    for value in values:
        bullet = cell.add_paragraph()
        style_paragraph(bullet)
        bullet.paragraph_format.left_indent = Inches(0.18)
        run = bullet.add_run("• " + value)
        style_run(run)


def _merged_section_row(table, text: str) -> None:
    row = table.add_row()
    cell = row.cells[0].merge(row.cells[1])
    _cell_text(cell, text, bold=True, clear=True)
    shade_cell(cell, SECTION_FILL)


def _flow_section(table, heading: str, sections: list[FlowSection]) -> None:
    _merged_section_row(table, heading)
    if not sections:
        row = table.add_row()
        cell = row.cells[0].merge(row.cells[1])
        _cell_text(cell, "None.", clear=True)
        return
    for section in sections:
        _merged_text_row(table, section.title, "")
        _flow_rows(table, section.items)


def _flow_rows(table, items: list[FlowItem]) -> None:
    for item in items:
        row = table.add_row()
        actor_cell, system_cell = row.cells
        target = system_cell if item.is_system else actor_cell
        _cell_text(
            target,
            f"{item.number}. {item.text}" if item.number else item.text,
            clear=True,
        )
        _cell_text(actor_cell if item.is_system else system_cell, "", clear=True)


def _cell_text(
    cell,
    text: str,
    bold: bool = False,
    alignment=WD_ALIGN_PARAGRAPH.LEFT,
    color: str | None = None,
    clear: bool = False,
) -> None:
    if clear:
        cell.text = ""
    paragraph = cell.paragraphs[0]
    style_paragraph(paragraph, alignment)
    if text:
        run = paragraph.add_run(text)
        style_run(run, bold=bold, color=color)
